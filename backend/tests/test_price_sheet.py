"""The team's price sheet: template, import, and quotes priced from its tables.

Worked examples use the untouched template, whose example costs equal the placeholder
selling prices with a 0% margin, so its White Carrara Original Q 645 x 3000 mm piece
costs 5,670.00 MXN, as in test_quote.py.
"""

import importlib.util
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from backend.app.ai.base import ModelError
from backend.app.ai.suggester import validate_suggestions
from backend.app.catalogue import Catalogue
from backend.app.quote import CountertopRun, QuoteError, QuoteRequest, SplashRun, price_quote

REPO = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location("price_sheet", REPO / "catalog" / "kober" / "price_sheet.py")
ps = importlib.util.module_from_spec(spec)
spec.loader.exec_module(ps)

BASE = Catalogue()


def sheet(edit=None) -> dict:
    """Template workbook -> tabs, after `edit(workbook)` changes some cells."""
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "precios.xlsx"
        ps.write_template(path, BASE)
        if edit:
            wb = load_workbook(path)
            edit(wb)
            wb.save(path)
        return ps.read_sheet(str(path))


def set_cell(wb, tab: str, key_col: int, key: str, col: int, value) -> None:
    for row in wb[tab].iter_rows(min_row=2):
        if str(row[key_col].value) == key:
            row[col].value = value
            return
    raise KeyError(key)


def imported(edit=None) -> Catalogue:
    prices, _ = ps.build_prices(sheet(edit), BASE, "test")
    return Catalogue(prices=prices)


def carrara(runs, **kw):
    return QuoteRequest("diseno-white-carrara", "original_q", tuple(runs), **kw)


class ParsingTests(unittest.TestCase):
    def test_5670_point_50_with_a_thousands_comma_is_567050_centavos(self):
        self.assertEqual(ps.money_minor("5,670.50", "x", []), 567050)
        self.assertEqual(ps.money_minor("$1,200", "x", []), 120000)
        self.assertEqual(ps.money_minor("5670,5", "x", []), 567050)

    def test_a_word_where_a_price_goes_is_an_error_not_zero(self):
        errors: list[str] = []
        self.assertIsNone(ps.money_minor("pendiente", "Cubiertas row 3", errors))
        self.assertIn("Cubiertas row 3", errors[0])

    def test_16_percent_written_three_ways_is_1600_basis_points(self):
        self.assertEqual([ps.percent_bp(v, "x", []) for v in ("16", "16%", "16.0 %")], [1600, 1600, 1600])


class ImportTests(unittest.TestCase):
    def test_the_untouched_template_imports_as_a_draft_with_every_piece_priced(self):
        prices, warnings = ps.build_prices(sheet(), BASE, "test")
        self.assertEqual(prices["status"], "DRAFT")
        self.assertTrue(prices["version"].startswith("sheet-"))
        self.assertIn("645x3000", prices["countertop_piece_minor"]["estandar"])
        self.assertIn("645x1200", prices["countertop_piece_minor"]["premium_plus"])  # Estilo short piece
        self.assertNotIn("700x3000", prices["countertop_piece_minor"]["estandar"])  # not made over 645 mm
        self.assertTrue(any("margin" in w for w in warnings))

    def test_a_missing_piece_cost_stops_the_import_and_names_the_size(self):
        def blank(wb):
            for row in wb["Cubiertas"].iter_rows(min_row=2):
                if row[0].value == "estandar" and row[1].value == 645 and row[2].value == 3000:
                    row[3].value = None
        with self.assertRaises(ps.SheetError) as e:
            ps.build_prices(sheet(blank), BASE, "test")
        self.assertIn("estandar has no cost for 645x3000", str(e.exception))

    def test_an_unknown_category_and_a_bad_status_are_both_reported(self):
        def break_it(wb):
            set_cell(wb, "Ajustes", 0, "estado", 1, "listo")
            set_cell(wb, "Acabados", 0, "diseno-white-carrara", 4, "lujo")
        with self.assertRaises(ps.SheetError) as e:
            ps.build_prices(sheet(break_it), BASE, "test")
        self.assertIn("BORRADOR or VIGENTE", str(e.exception))
        self.assertIn("diseno-white-carrara: category 'lujo'", str(e.exception))

    def test_a_basik_finish_cannot_be_moved_to_a_premium_category(self):
        with self.assertRaises(ps.SheetError):
            ps.build_prices(sheet(lambda wb: set_cell(wb, "Acabados", 0, "basik-buka-bark", 4, "premium")), BASE, "t")


class SheetQuoteTests(unittest.TestCase):
    def test_run_2300mm_carrara_from_the_template_is_one_3000mm_piece_at_5670_mxn(self):
        q = price_quote(carrara([CountertopRun("A", 2300)]), imported())
        piece = next(ln for ln in q.lines if ln.kind == "countertop_piece")
        self.assertEqual((piece.qty, piece.unit_price.minor), (1, 567000))
        self.assertEqual(q.price_list_status, "DRAFT")
        self.assertIn("in review", q.assumptions[0])

    def test_a_25_percent_estandar_margin_makes_the_5670_piece_7087_50(self):
        cat = imported(lambda wb: set_cell(wb, "Margenes", 0, "estandar", 1, "25"))
        piece = next(ln for ln in price_quote(carrara([CountertopRun("A", 2300)]), cat).lines
                     if ln.kind == "countertop_piece")
        self.assertEqual(piece.unit_price.minor, 708750)

    def test_essence_at_120_percent_of_original_q_makes_the_5670_piece_6804(self):
        cat = imported(lambda wb: set_cell(wb, "Perfiles", 0, "essence", 2, "120"))
        req = QuoteRequest("diseno-white-carrara", "essence", (CountertopRun("A", 2300),))
        piece = next(ln for ln in price_quote(req, cat).lines if ln.kind == "countertop_piece")
        self.assertEqual(piece.unit_price.minor, 680400)

    def test_sheet_prices_with_iva_included_are_stored_without_it_so_5670_becomes_4887_93(self):
        cat = imported(lambda wb: set_cell(wb, "Ajustes", 0, "precios_incluyen_iva", 1, "si"))
        piece = next(ln for ln in price_quote(carrara([CountertopRun("A", 2300)]), cat).lines
                     if ln.kind == "countertop_piece")
        self.assertEqual(piece.unit_price.minor, 488793)  # 567000 / 1.16, rounded half up

    def test_splash_3100_by_600_caracatta_panel_with_10_percent_margin_is_3326_40(self):
        cat = imported(lambda wb: set_cell(wb, "Margenes", 0, "premium_plus", 2, "10"))
        req = carrara([CountertopRun("A", 2300)], splash_finish_id="estilo-caracatta",
                      splash_runs=(SplashRun("S1", 3100, 600),))
        panel = next(ln for ln in price_quote(req, cat).lines if ln.kind == "splash_panel")
        self.assertEqual((panel.qty, panel.unit_price.minor), (1, 332640))

    def test_a_finish_moved_to_premium_is_priced_from_the_premium_table(self):
        cat = imported(lambda wb: set_cell(wb, "Acabados", 0, "diseno-white-carrara", 4, "premium"))
        self.assertEqual(cat.finish("diseno-white-carrara").category, "premium")
        piece = next(ln for ln in price_quote(carrara([CountertopRun("A", 2300)]), cat).lines
                     if ln.kind == "countertop_piece")
        self.assertEqual(piece.unit_price.minor, 756000)  # 2,400 MXN/m x 3.0 m x Original Q (template)

    def test_a_live_sheet_drops_the_review_note_and_confirmed_iva_says_included(self):
        def live(wb):
            set_cell(wb, "Ajustes", 0, "estado", 1, "VIGENTE")
            set_cell(wb, "Ajustes", 0, "iva_confirmado", 1, "si")
        q = price_quote(carrara([CountertopRun("A", 2300)], language="es"), imported(live))
        self.assertEqual(q.price_list_status, "LIVE")
        self.assertFalse(any("revisión" in a or "ejemplo" in a for a in q.assumptions))
        self.assertIn("IVA 16% incluido en el total.", q.assumptions)

    def test_a_piece_the_sheet_does_not_price_is_refused_not_guessed(self):
        prices, _ = ps.build_prices(sheet(), BASE, "test")
        del prices["countertop_piece_minor"]["estandar"]["645x3000"]
        with self.assertRaises(QuoteError):
            price_quote(carrara([CountertopRun("A", 2300)]), Catalogue(prices=prices))


class AvailabilityTests(unittest.TestCase):
    def test_a_finish_marked_not_available_is_never_suggested(self):
        cat = imported(lambda wb: set_cell(wb, "Acabados", 0, "estilo-caracatta", 5, "no"))
        good = {"title": "x", "countertop_finish_id": "estilo-caracatta", "profile_id": "original_q",
                "backsplash_finish_id": "none", "reason": "x"}
        with self.assertRaises(ModelError):
            validate_suggestions({"designs": [good]}, cat)
        self.assertFalse(cat.available("estilo-caracatta"))


if __name__ == "__main__":
    unittest.main()
